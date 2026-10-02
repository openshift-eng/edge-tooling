#!/usr/bin/env python3
"""Assemble the release-planning report from checks.json and recommendations.json.

Layout (see references/release-planning-report-template.md):

    verdict → decisions → cut line → people over target → dormant scope →
    process gaps → how the numbers are computed → collapsible appendix

Both .md and .docx are rendered from the same block model so they cannot
drift. Every figure comes from checks.json; recommendations.json supplies
narrative only. Feature keys are rendered as "Name (KEY)" and linked.
"""

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(__file__))
from _jira_transforms import load_json

JIRA_BASE = "https://redhat.atlassian.net/browse"
JIRA_KEY_RE = re.compile(r"(?<!\[)(?<!/)\b(OCPSTRAT-\d+|OCPEDGE-\d+|USHIFT-\d+|OCPBUGS-\d+)\b(?!\])")
EMOJI_RE = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\uFE0F]")
GENDERED_RE = re.compile(r"\b(he|she|him|her|his|hers|himself|herself)\b", re.IGNORECASE)

RISK_GLYPH = {"HIGH": "🔴", "MEDIUM": "🟡", "LOW": "🟢", "N/A": "⚪"}
FITS_GLYPH = {"fits": "✅", "partial": "⚠️ partial", "over": "❌"}
METHOD_DOC = "plugins/edge-scrum/references/release-planning-method.md"

# Section rendering order and H2 headings, shared by the DOCX and HTML renderers
# so they cannot drift. None = the block carries no heading of its own.
SECTION_ORDER = [
    ("VERDICT", None), ("SUMMARY_STRIP", None), ("STATS_LINE", None),
    ("DECISIONS", "Decisions needed this week"), ("CUT_LINE", "Where the cut line falls"),
    ("PEOPLE", "People over target"), ("DORMANT", "Scope nobody has started"),
    ("GAPS", "Process gaps"), ("METHOD", "How the numbers are computed"),
    ("APPENDIX", None), ("FOOTER", None),
]

# Risk token → CSS class for HTML risk cells (mirrors the DOCX RISK_COLORS palette).
RISK_CLASS = {
    "HIGH": "r-high", "OVER": "r-high", "FAIL": "r-high", "NO_CONTRIBUTORS": "r-high",
    "MEDIUM": "r-med", "WARN": "r-med", "NOT_IN_ROSTER": "r-med", "MISMATCH": "r-med",
    "LOW": "r-low", "OK": "r-low", "PASS": "r-low",
}


def short_version(v):
    """Strip a scheme prefix from a version tag: 'openshift-4.22' -> '4.22'."""
    return re.sub(r"^[A-Za-z]+-", "", str(v))


def version_major(v):
    """Major version number as a string, or None: 'openshift-4.22' -> '4'."""
    m = re.search(r"(\d+)\.\d+", str(v))
    return m.group(1) if m else None


def cut_fits_glyph(row):
    """Fits? cell for the cut line. A feature can fit within total team capacity yet
    still be at timeline risk because the people currently on it can't finish in time;
    mark those ✅⚠️ so the tick isn't read as 'all clear'."""
    if row["fits"] == "fits" and row.get("timeline_risk") in ("HIGH", "NO_CONTRIBUTORS"):
        return "✅⚠️"
    return FITS_GLYPH[row["fits"]]


def jira_linkify(text):
    return JIRA_KEY_RE.sub(rf"[\1]({JIRA_BASE}/\1)", text)


def fmt_sp(sp, unpointed=0):
    """'90 SP' or '90 SP +12 unptd' — pointed and unpointed never merged."""
    s = f"{sp} SP"
    if unpointed:
        s += f" +{unpointed} unptd"
    return s


def sme_or_dash(sme):
    return "—" if sme in (None, "", "None") else sme


def plural(n, noun, plural_noun=None):
    try:
        n_int = int(n)
    except (TypeError, ValueError):
        return f"{n} {plural_noun or noun + 's'}"
    return f"{n} {noun if n_int == 1 else (plural_noun or noun + 's')}"


# --- Block model -----------------------------------------------------------
# Each block is a dict: {"t": "h2"|"p"|"table"|"bullets"|"details"|"small", ...}


def feature_label(names, key):
    name = names.get(key)
    return f"{name} ({key})" if name else key


def build_blocks(checks, recs, params):
    meta = checks["meta"]
    names = checks.get("feature_names", {})
    blocks = {}

    risk = meta.get("overall_risk", "LOW")
    gap = meta.get("gap_sp", 0)
    gap_hidden_low = meta.get("gap_with_hidden_low", gap)
    headline = recs.get("headline") or recs.get("executive_summary") or ""
    glyph = RISK_GLYPH.get(risk, "")
    if gap < 0:
        verdict = f"**{glyph} {risk} — ~{abs(gap)} SP must leave the release.**"
    elif gap_hidden_low < 0:
        verdict = (f"**{glyph} {risk} — pointed scope fits with {gap} SP to spare, but not once unpointed work "
                   f"is estimated ({gap_hidden_low:+d} SP at the high estimate).**")
    else:
        verdict = f"**{glyph} {risk} — scope fits with {gap} SP to spare.**"
    verdict += (f" {plural(params['remaining_sprints'], 'dev sprint')} left · pencils down S{params['pencils_down']}"
                f" · assessed {params['today']}")
    # Version-mismatch caveat: epics tagged for another release are excluded from the
    # math, but open work under them may really belong to this release. Surface it in
    # the verdict rather than aliasing versions — the fix belongs in Jira.
    mistagged = [e for e in checks.get("process_gaps", {}).get("excluded_epics", [])
                 if str(e.get("reason", "")).startswith("epic targets") and e.get("open_sp", 0) > 0]
    if mistagged:
        mis_sp = sum(e["open_sp"] for e in mistagged)
        # Group open SP by the version each epic is tagged for; lumping them together
        # over-states the question, because an adjacent same-major release (5.0 carry-over,
        # 5.2 future) is plausibly correct while a different major (a MicroShift 4.x tag
        # under a 5.x release) is the genuinely ambiguous one to confirm.
        by_version = {}
        for e in mistagged:
            v = str(e["reason"])[len("epic targets "):].strip() or "unknown"
            by_version[v] = by_version.get(v, 0) + e["open_sp"]
        groups = " · ".join(f"{sp} SP {v}" for v, sp in sorted(by_version.items()))
        target_major = str(params["version"]).split(".")[0]
        suspect = sorted({short_version(v) for v in by_version if version_major(v) not in (None, target_major)})
        verdict += (f". A further {mis_sp} SP under {plural(len(mistagged), 'epic')} is tagged for "
                    f"other releases ({groups})")
        verdict += f"; confirm the {', '.join(suspect)} tags." if suspect else "."
    if str(params["remaining_sprints"]) == "0":
        verdict += "\n\n**Note:** 0 sprints remain before pencils down, so capacity is zero and every open story point is overrun. Read this as a post-cut inventory, not a forecast."
    if headline:
        verdict += f"\n\n{headline}"
    blocks["VERDICT"] = [{"t": "p", "text": verdict}]

    hidden = checks.get("hidden_scope", {})
    blocks["SUMMARY_STRIP"] = [{
        "t": "table",
        "headers": ["Scope (pointed)", "Hidden scope", "Capacity", "Gap"],
        "align": ["r", "r", "r", "r"],
        "rows": [[
            f"**{meta['total_remaining_sp']} SP**",
            f"{hidden.get('unpointed_open', 0)} stories unpointed ({hidden.get('unpointed_assigned', 0)} assigned)"
            f" ≈ {hidden.get('estimate_low', 0)}–{hidden.get('estimate_high', 0)} SP",
            f"**{meta['total_capacity_sp']} SP**",
            f"**{gap:+d} SP** ({meta.get('gap_with_hidden_low', gap):+d} to {meta.get('gap_with_hidden_high', gap):+d} incl. hidden)",
        ]],
    }]

    blocks["STATS_LINE"] = [{"t": "small", "text": (
        f"{plural(meta['active_features'], 'active feature')} · {meta['dormant_features']} dormant · "
        f"{plural(meta['overloaded_people_count'], 'person', 'people')} over target · "
        f"{plural(meta['spof_features_count'], 'single-owner feature')} · "
        f"{meta['high_risk_count']} HIGH / {meta['medium_risk_count']} MEDIUM / {meta['low_risk_count']} LOW"
    )}]

    # Decisions
    timeline_map = {t["feature_key"]: t for t in checks.get("timeline", [])}
    rows = []
    for i, d in enumerate(recs.get("decisions", [])[:5], 1):
        text = d.get("decision", "")
        why = d.get("why", "")
        keys_in_text = list(dict.fromkeys(JIRA_KEY_RE.findall(text)))
        if not why and len(keys_in_text) == 1:
            # Attach the derivation only when the decision is about one feature
            # whose timeline projection is the reason for it.
            t = timeline_map.get(keys_in_text[0])
            if t and t.get("risk") == "NO_CONTRIBUTORS":
                why = f"{t['remaining_sp']} SP open with no roster contributor assigned"
            elif t and t.get("risk") == "HIGH":
                why = (f"{t['remaining_sp']} SP ÷ {t['velocity_per_sprint']} SP/sprint = "
                       f"{t['sprints_needed']} sprints vs {t['sprints_left']} left")
        if why:
            text += f" <sub>Why: {why}</sub>"
        rows.append([str(i), text, d.get("frees", ""), d.get("owner", ""), d.get("by", "")])
    if rows:
        blocks["DECISIONS"] = [{"t": "table", "headers": ["#", "Decision", "Frees", "Owner", "By"],
                                "align": ["l", "l", "r", "l", "l"], "rows": rows}]
    else:
        blocks["DECISIONS"] = [{"t": "p", "text": "*No decisions proposed.*"}]

    # Cut line
    cut = checks.get("cut_line", [])
    rows = []
    for r in cut:
        rows.append([str(r["rank"]) if r["rank"] is not None else "—",
                     feature_label(names, r["feature_key"]), sme_or_dash(r["sme"]),
                     fmt_sp(r["remaining_sp"], r.get("unpointed_open", 0)),
                     str(r["cumulative_sp"]), cut_fits_glyph(r)])
    blocks["CUT_LINE"] = [
        {"t": "p", "text": f"Active features in PM rank order; cumulative pointed SP against {meta['total_capacity_sp']} SP capacity."},
        {"t": "table", "headers": ["Rank", "Feature", "SME", "Left", "Cum.", "Fits?"],
         "align": ["r", "l", "l", "r", "r", "c"], "rows": rows},
        {"t": "small", "text": "✅⚠️ = fits within total team capacity but not at the current allocation of people (see timeline). "
                              "Rank comes from Jira; if the stack rank is stale, so is this table. Unpointed work is shown but not summed."},
    ]

    # People over target
    cap = checks.get("capacity", [])
    over = [c for c in cap if c["status"] == "OVER"]
    non_roster = [c for c in cap if not c["in_roster"] and (c["assigned_sp"] or c["unpointed_assigned"])]
    rows = []
    for c in over + non_roster:
        heaviest = ""
        best = 0
        for fk in c["features"]:
            t = timeline_map.get(fk)
            if t and t["remaining_sp"] > best:
                best, heaviest = t["remaining_sp"], fk
        spof_for = {s["feature_key"] for s in checks["assignment"]["spof"] if s["sole_contributor"] == c["person"]}
        heaviest_txt = (feature_label(names, heaviest) if heaviest else "—") + (" (sole owner)" if heaviest in spof_for else "")
        target = "not on roster" if not c["in_roster"] else str(c["remaining_capacity"])
        rows.append([c["display_name"], f"**{c['assigned_sp']} SP**", target,
                     str(c["unpointed_assigned"]), heaviest_txt])
    hidden_heavy = [c for c in cap if c["status"] == "OK" and c["unpointed_assigned"] >= 3]
    people_blocks = []
    if rows:
        people_blocks.append({"t": "table", "headers": ["Person", "Assigned", "Capacity", "Unpointed", "Heaviest feature"],
                              "align": ["l", "r", "r", "r", "l"], "rows": rows})
    else:
        people_blocks.append({"t": "p", "text": "*Nobody is over target.*"})
    notes = []
    if hidden_heavy:
        notes.append(", ".join(f"{c['display_name']} ({c['unpointed_assigned']} unpointed)" for c in hidden_heavy)
                     + " look OK on points but carry unpointed work and are likely full.")
    if non_roster:
        notes.append(f"{plural(len(non_roster), 'assignee')} missing from `.roster.json` — capacity totals are unreliable until reconciled.")
    if notes:
        people_blocks.append({"t": "small", "text": " ".join(notes)})
    blocks["PEOPLE"] = people_blocks

    # Dormant
    dormant = checks.get("scope_confirmation", [])
    if dormant:
        keys = ", ".join(d["feature_key"] for d in dormant) if len(dormant) <= 8 else "listed below"
        no_sme = sum(1 for d in dormant if d["sme"] in (None, "", "None"))
        n_dormant = len(dormant)
        if n_dormant == 1:
            lead = f"1 feature is still **New** with no evidence of work in {params['version']}: {keys}."
            sme_note = " It has no SME." if no_sme else ""
        else:
            lead = f"{n_dormant} features are still **New** with no evidence of work in {params['version']}: {keys}."
            sme_note = f" {no_sme} of them have no SME." if no_sme else ""
        text = (lead + sme_note
                + " Confirm as *deferred* or *committed*; anything committed needs epics and stories before the next planning.")
        d_blocks = [{"t": "p", "text": text}]
        for line in recs.get("scope_decisions", [])[:3]:
            d_blocks.append({"t": "bullets", "items": [line]})
        d_blocks.append({"t": "details", "summary": "Dormant features and why",
                         "blocks": [{"t": "table", "headers": ["Feature", "SME", "Epics", "Stories", "Why dormant"],
                                     "rows": [[feature_label(names, d["feature_key"]), sme_or_dash(d["sme"]),
                                               str(d["epic_count"]), str(d["story_count"]), d["reason"]] for d in dormant]}]})
        blocks["DORMANT"] = d_blocks
    else:
        blocks["DORMANT"] = [{"t": "p", "text": "*Every targeted feature shows activity.*"}]

    # Process gaps: computed facts first, narrative after
    pg = checks.get("process_gaps", {})
    items = []
    drift = pg.get("roster_drift", [])
    if drift:
        items.append(f"**Roster drift** — {plural(len(drift), 'person', 'people')} carry work but aren't in the roster ("
                     + ", ".join(d["display_name"] for d in drift) + "). Cheap fix, large effect on every capacity number.")
    if pg.get("unpointed_open_stories"):
        items.append(f"**Just-in-time pointing** — {pg['unpointed_open_stories']} open stories unpointed "
                     f"({pg.get('unpointed_assigned_stories', 0)} already assigned); the {meta['total_remaining_sp']} SP figure is a floor.")
    ns = pg.get("features_without_sme", {})
    total_no_sme = len(ns.get("active", [])) + len(ns.get("dormant", []))
    if total_no_sme:
        items.append(f"**No SME on {total_no_sme} of {meta['total_features']} features** — "
                     f"{len(ns.get('active', []))} active ({', '.join(ns.get('active', [])) or '—'}), {len(ns.get('dormant', []))} dormant.")
    if pg.get("sized_active_features") and pg.get("undersized_features"):
        undersized = pg["undersized_features"]
        verb, poss = ("is", "its") if int(undersized) == 1 else ("are", "their")
        items.append(f"**Sizing** — {undersized} of {pg['sized_active_features']} sized active features "
                     f"{verb} larger than {poss} T-shirt implies.")
    if pg.get("excluded_epics"):
        items.append(f"**Out-of-release work under {params['version']} features** — {plural(len(pg['excluded_epics']), 'epic')} ({pg.get('excluded_open_sp', 0)} open SP) "
                     "are Dev Complete/Closed or target another release and were left out of the math.")
    # Features whose only epics were excluded read as "no epics created" if left alone;
    # they are finished or mis-targeted, a distinct condition from un-refined scope.
    excluded_only = [d for d in checks.get("data_quality", []) if d.get("all_epics_excluded")]
    if excluded_only:
        n_excl = len(excluded_only)
        has_verb = "has" if n_excl == 1 else "have"
        subject = "it is" if n_excl == 1 else "they are"
        items.append(f"**{plural(n_excl, 'feature')} {has_verb} only done or out-of-release epics** — "
                     + ", ".join(d["feature_key"] for d in excluded_only)
                     + f"; {subject} finished or mis-targeted, not un-refined.")
    items.extend(recs.get("process_gaps", [])[:4])
    blocks["GAPS"] = [{"t": "bullets", "items": items}] if items else [{"t": "p", "text": "*No systemic gaps detected.*"}]

    # Method
    rows = [[m["label"], m["formula"], m["shown_as"]] for m in checks.get("method", [])]
    limits = [
        "timeline and capacity signals both depend on contributor load, so an overloaded person raises the risk of every feature they touch",
        "capacity ignores PTO, holidays and bug-fix time — treat it as an upper bound",
    ]
    blocks["METHOD"] = [
        {"t": "table", "headers": ["Figure", "Formula", "This run"], "align": ["l", "l", "l"], "rows": rows},
        {"t": "small", "text": f"Full method, thresholds and known limitations: `{METHOD_DOC}`. Known limitations this run: " + "; ".join(limits) + "."},
    ]

    # Appendix
    app = []
    comp_rows = [[feature_label(names, c["feature_key"]), f"{RISK_GLYPH.get(c['composite_risk'], '')} {c['composite_risk']}",
                  c["data_quality"], c["timeline"], c["capacity"], c["assignment"], c["bugs"], c["sizing"]]
                 for c in checks.get("composite", [])]
    app.append({"t": "h3", "text": "Composite risk by feature"})
    app.append({"t": "table", "headers": ["Feature", "Risk", "Data quality", "Timeline", "Capacity", "Assignment", "Bugs", "Sizing"],
                "rows": comp_rows, "risk_cols": {1, 2, 3, 4, 5, 6, 7}})

    tl_rows = [[feature_label(names, t["feature_key"]), fmt_sp(t["remaining_sp"], t.get("unpointed_open", 0)),
                str(t["velocity_per_sprint"]), str(t["sprints_needed"]), str(t["gap"]), t["risk"]]
               for t in checks.get("timeline", []) if t["risk"] != "N/A"]
    app.append({"t": "h3", "text": "Timeline projection"})
    app.append({"t": "table", "headers": ["Feature", "Left", "Velocity/sprint", "Sprints needed", "Gap", "Risk"],
                "rows": tl_rows, "risk_cols": {5}})

    cap_rows = [[c["display_name"], str(c["assigned_sp"]), str(c["unpointed_assigned"]), str(c["remaining_capacity"]),
                 c["status"], ", ".join(c["features"])] for c in cap]
    app.append({"t": "h3", "text": "Capacity by person"})
    app.append({"t": "table", "headers": ["Person", "Assigned SP", "Unpointed", "Capacity", "Status", "Features"],
                "rows": cap_rows, "risk_cols": {4}})

    dq_rows = [[feature_label(names, d["feature_key"]), str(d["epic_count"]), str(d["story_count"]), f"{d['pointed_pct']}%",
                d["status"], d["reason"]] for d in checks.get("data_quality", [])]
    app.append({"t": "h3", "text": "Data quality"})
    app.append({"t": "table", "headers": ["Feature", "Epics", "Stories", "Pointed", "Status", "Note"], "rows": dq_rows, "risk_cols": {4}})

    spof_rows = [[feature_label(names, s["feature_key"]), s["sole_contributor_display"]] for s in checks["assignment"]["spof"]]
    ua_rows = [[feature_label(names, u["feature_key"]), str(u["count"]), str(u["sp"])] for u in checks["assignment"]["unassigned"]]
    app.append({"t": "h3", "text": "Assignment"})
    app.append({"t": "table", "headers": ["Single-owner feature", "Sole contributor"], "rows": spof_rows})
    app.append({"t": "table", "headers": ["Feature with unassigned stories", "Stories", "SP"], "rows": ua_rows})

    bl = checks.get("bug_load", {})
    bc_rows = [[b["key"], b["priority"], b["component"], b["status"]] for b in bl.get("unassigned_blocker_critical", [])]
    comp_bug_rows = [[comp, str(d["total"]), str(d["blocker"]), str(d["critical"]), str(d["unassigned"])]
                     for comp, d in bl.get("by_component", {}).items()]
    app.append({"t": "h3", "text": "Bug load"})
    app.append({"t": "table", "headers": ["Unassigned Blocker/Critical", "Priority", "Component", "Status"], "rows": bc_rows})
    app.append({"t": "table", "headers": ["Component", "Total", "Blocker", "Critical", "Unassigned"], "rows": comp_bug_rows})

    sz_rows = [[feature_label(names, s["feature_key"]), s["tshirt"], str(s.get("total_sp", s["actual_sp"])), str(s["epic_count"]),
                str(s["contributor_count"]), str(s.get("dedicated_sprints", "N/A")), s["assessment"]] for s in checks.get("sizing", [])]
    app.append({"t": "h3", "text": "Sizing (informational)"})
    app.append({"t": "table", "headers": ["Feature", "T-shirt", "Total SP", "Epics", "Contributors", "Dedicated sprints", "Assessment"],
                "rows": sz_rows})

    ex = pg.get("excluded_epics", [])
    if ex:
        app.append({"t": "h3", "text": "Epics excluded from scope"})
        app.append({"t": "table", "headers": ["Feature", "Epic", "Status", "Open SP", "Why"],
                    "rows": [[feature_label(names, e["feature_key"]), e["key"], e["status"], str(e["open_sp"]), e["reason"]] for e in ex]})

    if recs.get("per_feature") or recs.get("per_person"):
        app.append({"t": "h3", "text": "Detailed recommendations"})
        if recs.get("per_feature"):
            app.append({"t": "bullets", "items": recs["per_feature"]})
        if recs.get("per_person"):
            app.append({"t": "bullets", "items": recs["per_person"]})

    blocks["APPENDIX"] = [{"t": "details", "summary": "Appendix — full check results", "blocks": app, "page_break": True}]

    cmd = (f"/release-planning {params['version']} {params['first_sprint']}-{params['last_sprint']} "
           f"bc:{params['last_sprint']} pd:{params['pencils_down']}")
    blocks["FOOTER"] = [{"t": "small", "text": f"Generated by `{cmd}` · read-only against Jira · "
                                               f"component filter: {meta.get('component_filter', 'none')}"}]
    return blocks


# --- Markdown renderer -----------------------------------------------------


def _md_table(headers, rows, align=None):
    align = align or ["l"] * len(headers)
    sep = {"l": "---", "r": "---:", "c": ":---:"}
    lines = ["| " + " | ".join(headers) + " |",
             "| " + " | ".join(sep.get(a, "---") for a in align) + " |"]
    if not rows:
        lines.append("| " + " | ".join(["*none*"] + [""] * (len(headers) - 1)) + " |")
    for row in rows:
        cells = [str(c).replace("|", "\\|").replace("\n", " ") for c in row]
        while len(cells) < len(headers):
            cells.append("")
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def render_blocks_md(blocks):
    out = []
    for b in blocks:
        t = b["t"]
        if t == "p":
            out.append(b["text"])
        elif t == "small":
            out.append(f"<sub>{b['text']}</sub>")
        elif t == "h3":
            out.append(f"### {b['text']}")
        elif t == "table":
            out.append(_md_table(b["headers"], b["rows"], b.get("align")))
        elif t == "bullets":
            out.append("\n".join(f"- {i}" for i in b["items"]))
        elif t == "details":
            inner = render_blocks_md(b["blocks"])
            out.append(f"<details>\n<summary><b>{b['summary']}</b></summary>\n\n{inner}\n\n</details>")
    return "\n\n".join(out)


def render_markdown(checks, recs, template, params):
    blocks = build_blocks(checks, recs, params)
    text = template.replace("{VERSION}", str(params["version"]))
    for key, blist in blocks.items():
        text = text.replace(f"{{{key}}}", render_blocks_md(blist))
    return jira_linkify(text).rstrip() + "\n"


# --- DOCX renderer ---------------------------------------------------------


def render_docx(checks, recs, params, output_path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml

    DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
    HEADING_BLUE = RGBColor(0x2C, 0x5F, 0x8A)
    RISK_COLORS = {
        "HIGH": ("FADBD8", RGBColor(0xC0, 0x39, 0x2B)), "OVER": ("FADBD8", RGBColor(0xC0, 0x39, 0x2B)),
        "FAIL": ("FADBD8", RGBColor(0xC0, 0x39, 0x2B)), "NO_CONTRIBUTORS": ("FADBD8", RGBColor(0xC0, 0x39, 0x2B)),
        "MEDIUM": ("FEF9E7", RGBColor(0xB7, 0x95, 0x0B)), "WARN": ("FEF9E7", RGBColor(0xB7, 0x95, 0x0B)),
        "NOT_IN_ROSTER": ("FEF9E7", RGBColor(0xB7, 0x95, 0x0B)), "MISMATCH": ("FEF9E7", RGBColor(0xB7, 0x95, 0x0B)),
        "LOW": ("D5F5E3", RGBColor(0x1E, 0x8E, 0x3E)), "OK": ("D5F5E3", RGBColor(0x1E, 0x8E, 0x3E)),
        "PASS": ("D5F5E3", RGBColor(0x1E, 0x8E, 0x3E)),
    }
    ALT_ROW_BG = "F2F2F2"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for level, sz, color in [("Heading 1", 20, DARK_BLUE), ("Heading 2", 15, HEADING_BLUE), ("Heading 3", 12, HEADING_BLUE)]:
        s = doc.styles[level]
        s.font.name = "Calibri"
        s.font.size = Pt(sz)
        s.font.color.rgb = color
        s.font.bold = True
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2.0)
        section.left_margin = section.right_margin = Cm(2.0)

    def clean(text):
        text = EMOJI_RE.sub("", str(text))
        text = re.sub(r"<sub>(.*?)</sub>", r" (\1)", text)
        text = re.sub(r"<[^>]+>", "", text)
        text = text.replace("**", "").replace("`", "")
        text = re.sub(r"(?<!\w)[_*](.+?)[_*](?!\w)", r"\1", text)
        text = JIRA_KEY_RE.sub(r"\1", text)
        return text.strip()

    def set_shading(cell, color_hex):
        cell._element.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>'))

    def add_table(headers, rows, risk_cols=None):
        risk_cols = risk_cols or set()
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(clean(h))
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_shading(cell, "2C5F8A")
        if not rows:
            rows = [["none"] + [""] * (len(headers) - 1)]
        for ri, row in enumerate(rows):
            cells = table.add_row().cells
            for ci, val in enumerate(row[:len(headers)]):
                cell = cells[ci]
                cell.text = ""
                text = clean(val)
                run = cell.paragraphs[0].add_run(text)
                run.font.size = Pt(9)
                token = text.strip().upper().split(" ")[0] if text else ""
                if ci in risk_cols and token in RISK_COLORS:
                    bg, fg = RISK_COLORS[token]
                    set_shading(cell, bg)
                    run.font.color.rgb = fg
                    run.bold = True
                elif ri % 2 == 1:
                    set_shading(cell, ALT_ROW_BG)
        tbl_pr = table._tbl.tblPr
        tbl_pr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>' + "".join(
                f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
                for side in ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>"))
        doc.add_paragraph()

    def render(blist):
        for b in blist:
            t = b["t"]
            if t == "p":
                p = doc.add_paragraph()
                text = clean(b["text"])
                bold = str(b["text"]).startswith("**")
                first, _, rest = text.partition(".")
                if bold and rest:
                    r = p.add_run(first + ".")
                    r.bold = True
                    p.add_run(rest)
                else:
                    p.add_run(text)
            elif t == "small":
                p = doc.add_paragraph()
                r = p.add_run(clean(b["text"]))
                r.font.size = Pt(8.5)
                r.italic = True
            elif t == "h3":
                doc.add_heading(clean(b["text"]), level=3)
            elif t == "table":
                add_table(b["headers"], b["rows"], b.get("risk_cols"))
            elif t == "bullets":
                for item in b["items"]:
                    doc.add_paragraph(clean(item), style="List Bullet")
            elif t == "details":
                if b.get("page_break"):
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                    doc.add_heading(clean(b["summary"]), level=2)
                else:
                    doc.add_heading(clean(b["summary"]), level=3)
                render(b["blocks"])

    blocks = build_blocks(checks, recs, params)
    doc.add_heading(f"OCP {params['version']} Planning Risk", level=1)
    for key, heading in SECTION_ORDER:
        if heading:
            doc.add_heading(heading, level=2)
        render(blocks.get(key, []))

    doc.save(output_path)
    print(f"Wrote {output_path}", file=sys.stderr)


# --- HTML renderer ---------------------------------------------------------

_HTML_ESCAPE = (("&", "&amp;"), ("<", "&lt;"), (">", "&gt;"))
_ALIGN_CSS = {"l": "left", "r": "right", "c": "center"}


def _html_escape(text):
    """Entity-encode HTML metacharacters for a plain text context (no markup)."""
    text = str(text)
    for a, b in _HTML_ESCAPE:
        text = text.replace(a, b)
    return text


def _html_inline(text):
    """Convert the report's inline markup to HTML, escaping metacharacters.

    Order matters: <sub> is preserved verbatim across escaping; then `code`,
    **bold**, *italic* and bare Jira keys become tags. Escaping runs before the
    markup substitutions so user text can never inject markup.
    """
    text = str(text)
    text = text.replace("<sub>", "\x00SUB\x00").replace("</sub>", "\x00/SUB\x00")
    for a, b in _HTML_ESCAPE:
        text = text.replace(a, b)
    text = text.replace("\x00SUB\x00", "<sub>").replace("\x00/SUB\x00", "</sub>")
    text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"(?<!\w)\*(.+?)\*(?!\w)", r"<em>\1</em>", text)
    text = JIRA_KEY_RE.sub(rf'<a href="{JIRA_BASE}/\1">\1</a>', text)
    return text


def _html_table(headers, rows, align=None, risk_cols=None):
    align = align or ["l"] * len(headers)
    risk_cols = risk_cols or set()

    def style(i):
        a = _ALIGN_CSS.get(align[i] if i < len(align) else "l", "left")
        return f' style="text-align:{a}"' if a != "left" else ""

    out = ["<table>", "<thead>", "<tr>"]
    for i, h in enumerate(headers):
        out.append(f"<th{style(i)}>{_html_inline(h)}</th>")
    out += ["</tr>", "</thead>", "<tbody>"]
    if not rows:
        rows = [["none"] + [""] * (len(headers) - 1)]
    for row in rows:
        cells = [c for c in row[:len(headers)]]
        while len(cells) < len(headers):
            cells.append("")
        out.append("<tr>")
        for ci, val in enumerate(cells):
            cls = ""
            if ci in risk_cols:
                token = EMOJI_RE.sub("", str(val)).strip().upper().split(" ")[0]
                if token in RISK_CLASS:
                    cls = f' class="{RISK_CLASS[token]}"'
            out.append(f"<td{style(ci)}{cls}>{_html_inline(val)}</td>")
        out.append("</tr>")
    out += ["</tbody>", "</table>"]
    return "\n".join(out)


def render_blocks_html(blocks):
    out = []
    for b in blocks:
        t = b["t"]
        if t == "p":
            for para in str(b["text"]).split("\n\n"):
                out.append(f"<p>{_html_inline(para)}</p>")
        elif t == "small":
            out.append(f'<p class="note">{_html_inline(b["text"])}</p>')
        elif t == "h3":
            out.append(f"<h3>{_html_inline(b['text'])}</h3>")
        elif t == "table":
            out.append(_html_table(b["headers"], b["rows"], b.get("align"), b.get("risk_cols")))
        elif t == "bullets":
            items = "\n".join(f"<li>{_html_inline(i)}</li>" for i in b["items"])
            out.append(f"<ul>\n{items}\n</ul>")
        elif t == "details":
            inner = render_blocks_html(b["blocks"])
            out.append(
                f"<details>\n<summary><strong>{_html_inline(b['summary'])}</strong></summary>\n{inner}\n</details>")
    return "\n".join(out)


_HTML_STYLE = """
:root { color-scheme: light dark; }
* { box-sizing: border-box; }
body {
  font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Helvetica, Arial, sans-serif;
  max-width: 900px; margin: 2rem auto; padding: 0 1rem; line-height: 1.5;
  color: #1f2328; background: #ffffff;
}
h1 { font-size: 1.8rem; border-bottom: 1px solid #d0d7de; padding-bottom: .3em; }
h2 { font-size: 1.4rem; border-bottom: 1px solid #d0d7de; padding-bottom: .3em; margin-top: 1.8em; }
h3 { font-size: 1.1rem; }
p.note { color: #656d76; font-size: .85rem; }
code { background: #eff1f3; padding: .15em .3em; border-radius: 4px; font-size: .9em; }
a { color: #0969da; text-decoration: none; }
a:hover { text-decoration: underline; }
table { border-collapse: collapse; width: 100%; margin: 1em 0; font-size: .9rem; }
th, td { border: 1px solid #d0d7de; padding: 6px 13px; }
th { background: #f6f8fa; text-align: left; }
tbody tr:nth-child(2n) { background: #f6f8fa; }
/* Backgrounds keep the DOCX palette; text darkened to clear WCAG AA 4.5:1
   (C0392B/B7950B/1E8E3E on these tints were 4.19/2.72/3.60) while keeping the hue. */
td.r-high { background: #FADBD8; color: #A5281B; font-weight: 600; }
td.r-med  { background: #FEF9E7; color: #7A5C00; font-weight: 600; }
td.r-low  { background: #D5F5E3; color: #10662B; font-weight: 600; }
details { margin: 1em 0; }
summary { cursor: pointer; }
@media (prefers-color-scheme: dark) {
  body { color: #e6edf3; background: #0d1117; }
  h1, h2 { border-bottom-color: #30363d; }
  p.note { color: #8b949e; }
  code { background: #161b22; }
  a { color: #4493f8; }
  th, td { border-color: #30363d; }
  th { background: #161b22; }
  tbody tr:nth-child(2n) { background: #161b22; }
  td.r-high { background: #3a1d1a; color: #f5a9a0; }
  td.r-med  { background: #33300f; color: #e6d16b; }
  td.r-low  { background: #12331f; color: #7ee2a8; }
}
"""

_HTML_DOC = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{TITLE}}</title>
<style>{{STYLE}}</style>
</head>
<body>
{{BODY}}
</body>
</html>
"""


def render_html(checks, recs, params, output_path):
    """Emit a self-contained .html from the same block model as .md/.docx."""
    blocks = build_blocks(checks, recs, params)
    # --version is arbitrary text; escape once and reuse for both the <title> and <h1>
    # so a value like "</title><script>..." can't inject markup into the report.
    title = _html_escape(f"OCP {params['version']} Planning Risk")
    parts = [f"<h1>{title}</h1>"]
    for key, heading in SECTION_ORDER:
        if heading:
            parts.append(f"<h2>{heading}</h2>")
        rendered = render_blocks_html(blocks.get(key, []))
        if rendered:
            parts.append(rendered)
    html = (_HTML_DOC
            .replace("{{TITLE}}", title)
            .replace("{{STYLE}}", _HTML_STYLE)
            .replace("{{BODY}}", "\n".join(parts)))
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"Wrote {output_path}", file=sys.stderr)


# --- Validation of the narrative -------------------------------------------


def validate_recs(recs, checks):
    """Return a list of problems in recommendations.json. Empty list = OK."""
    problems = []
    known_keys = set(checks.get("feature_names", {}))
    texts = []
    for k in ("headline", "executive_summary"):
        if recs.get(k):
            texts.append(recs[k])
    for d in recs.get("decisions", []):
        texts.extend(str(d.get(f, "")) for f in ("decision", "frees", "owner", "by", "why"))
    for k in ("scope_decisions", "process_gaps", "per_feature", "per_person", "team_level"):
        texts.extend(recs.get(k, []))
    joined = "\n".join(texts)
    if "](" in joined and "atlassian.net" in joined:
        problems.append("pre-built Jira links found — write bare keys, the assembler links them")
    for m in GENDERED_RE.finditer(joined):
        problems.append(f"gendered pronoun '{m.group(0)}' — use they/them or the person's name")
        break
    for key in set(JIRA_KEY_RE.findall(joined)):
        if key.startswith("OCPSTRAT-") and known_keys and key not in known_keys:
            problems.append(f"{key} is referenced but is not in this run's data")
    if len(recs.get("decisions", [])) > 5:
        problems.append("more than 5 decisions — keep the table to the five that matter")
    return problems


def soft_warnings(recs):
    """Non-fatal advisories — printed but never trigger --strict. Catches process-gap
    bullets that duplicate each other (same Jira key set and same numbers), which read
    as two findings when they are one."""
    warnings = []
    seen = {}
    for b in recs.get("process_gaps", []) or []:
        keys = frozenset(JIRA_KEY_RE.findall(b))
        nums = frozenset(re.findall(r"\d+", b))
        if not keys:
            continue
        sig = (keys, nums)
        if sig in seen:
            # Report only the shared Jira keys — an allowlisted, non-sensitive
            # identifier that lets the reader find the two bullets. Never echo the
            # bullet text or any values (numbers included) pulled from the free-form
            # narrative: process_gaps can carry names or other sensitive data and
            # this warning is printed to stderr / logs.
            warnings.append("two process_gaps bullets look duplicated (same Jira keys "
                            f"{sorted(keys)}); consolidate them into one")
        else:
            seen[sig] = b
    return warnings


# List-valued fields in recommendations.json. The file is LLM-written narrative, so an
# empty section can arrive as null rather than []. Normalize once after loading so every
# consumer sees a list: validate_recs/soft_warnings iterate these, and build_blocks slices
# them (e.g. decisions[:5], scope_decisions[:3], process_gaps[:4]) — None would raise.
RECS_LIST_FIELDS = (
    "decisions", "scope_decisions", "process_gaps",
    "per_feature", "per_person", "team_level",
)


def normalize_recs(recs):
    """Coerce null/missing list fields in recommendations.json to []. Returns recs."""
    for k in RECS_LIST_FIELDS:
        if recs.get(k) is None:
            recs[k] = []
    return recs


# --- Main ------------------------------------------------------------------


def main():
    parser = argparse.ArgumentParser(description="Assemble release planning report")
    parser.add_argument("--checks", required=True)
    parser.add_argument("--recommendations", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--version", required=True)
    parser.add_argument("--today", required=True)
    parser.add_argument("--first-sprint", required=True)
    parser.add_argument("--last-sprint", required=True)
    parser.add_argument("--pencils-down", default=None)
    parser.add_argument("--remaining-sprints", required=True)
    parser.add_argument("--total-dev-sprints", required=True)
    parser.add_argument("--output", required=True, help="Output base path (without extension)")
    parser.add_argument("--strict", action="store_true", help="Fail if recommendations.json violates the writing rules")
    parser.add_argument("--open", dest="open_report", action="store_true",
                        help="Open the generated .html in the default browser (omit when running headless)")
    args = parser.parse_args()

    # --version flows into all three outputs (the .md/.docx text is not HTML-escaped);
    # validate it at the trust boundary with an allow-list so no output can carry markup.
    if not re.fullmatch(r"\d+\.\d+(?:\.\d+|\.z)?", str(args.version)):
        parser.error(f"--version must look like 5.1, 5.1.0 or 5.1.z, got {args.version!r}")

    checks = load_json(args.checks)
    recs = normalize_recs(load_json(args.recommendations))
    with open(args.template) as f:
        template = f.read()

    problems = validate_recs(recs, checks)
    for p in problems:
        print(f"WARNING: recommendations.json: {p}", file=sys.stderr)
    for w in soft_warnings(recs):
        print(f"WARNING: recommendations.json: {w}", file=sys.stderr)
    if problems and args.strict:
        sys.exit(2)

    params = {
        "version": args.version, "today": args.today,
        "first_sprint": args.first_sprint, "last_sprint": args.last_sprint,
        "pencils_down": args.pencils_down or args.last_sprint,
        "remaining_sprints": args.remaining_sprints, "total_dev_sprints": args.total_dev_sprints,
    }

    parent = os.path.dirname(args.output)
    if parent:
        os.makedirs(parent, exist_ok=True)
    with open(args.output + ".md", "w") as f:
        f.write(render_markdown(checks, recs, template, params))
    print(f"Wrote {args.output}.md", file=sys.stderr)
    render_docx(checks, recs, params, args.output + ".docx")
    html_output = args.output + ".html"
    render_html(checks, recs, params, html_output)

    if args.open_report:
        # Best-effort: on a headless/remote session this silently does nothing (or
        # warns); never let it change the exit code, the report is already written.
        try:
            import webbrowser
            webbrowser.open("file://" + os.path.abspath(html_output))
        except Exception as e:  # noqa: BLE001 - opening a browser must never fail the run
            print(f"WARNING: could not open {html_output} in a browser: {e}", file=sys.stderr)


if __name__ == "__main__":
    main()
